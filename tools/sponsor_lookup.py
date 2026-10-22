#!/usr/bin/env python3
"""US visa-sponsor lookup built from public government filing data.

Answers two questions the evaluation framework asks of every US posting for a
candidate who needs sponsorship (04-job-evaluation.md, Sponsorship Gate and
Sponsor Strength):

  1. Does this employer have a record of sponsoring my visa class?
     (`lookup`, `top`) - from the Department of Labor's LCA disclosure files
     (one row per Labor Condition Application, tagged H-1B, E-3 Australian,
     H-1B1 Chile or H-1B1 Singapore) and, for H-1B, the USCIS H-1B Employer
     Data Hub (petition approvals and denials per employer and fiscal year).
  2. What does this posting say about sponsorship? (`check-posting`) - a
     deterministic phrase scan, so "no sponsorship" / "US citizens only" /
     "clearance required" is caught in code before any scoring tokens are spent.

Nothing is fetched. You download the public files once (see
tools/README_SPONSOR_TOOL.md), `import` aggregates them into sponsor_data.json
in the repo root, and every later call reads only that file.

Usage:
  python3 tools/sponsor_lookup.py import FILE [FILE ...] [--classes e3,h1b,h1b1] [--out PATH]
  python3 tools/sponsor_lookup.py info
  python3 tools/sponsor_lookup.py lookup "Company" --visa e3 [--title TEXT] [--years 3] [--min N] [--strong N] [--json]
  python3 tools/sponsor_lookup.py top --visa e3 [--state NY] [--title TEXT] [--years 3] [--limit 25] [--json]
  python3 tools/sponsor_lookup.py check-posting [--in FILE] --visa e3 --needs-sponsorship yes \
      [--no-sponsorship drop|flag] [--has-clearance] [--json]

Exit 0 on success, 1 on a usage or data error.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import statistics
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA_FILE = ROOT / "sponsor_data.json"
DATA_VERSION = 1

# Visa classes with public employer-level filing data. LCA values are matched
# after lower-casing and collapsing whitespace; the one-letter codes are the
# pre-FY2017 file layout (R = H-1B, A = E-3 Australian, C/S = H-1B1).
VISA_CLASSES = {
    "e3": {"label": "E-3 (Australian citizens)", "lca": {"e-3 australian", "e-3", "e3"}, "codes": {"a"}},
    "h1b": {"label": "H-1B", "lca": {"h-1b", "h1b"}, "codes": {"r"}},
    "h1b1": {"label": "H-1B1 (Chile / Singapore)", "lca": {"h-1b1 chile", "h-1b1 singapore"}, "codes": {"c", "s"}},
}
# Visa paths with no employer-level public filing data: lookups say so instead
# of reporting "None found", which would read as a negative signal.
NO_DATA_PATHS = {
    "tn": "TN status is granted at the border or by USCIS without an LCA; no employer filing data is published.",
    "o1": "O-1 petitions are not published per employer.",
    "l1": "L-1 is an intra-company transfer from your current employer's foreign office; no public employer data.",
    "none": "No sponsorship needed - sponsor data does not apply.",
}

# Default "Strong" thresholds: filings in the lookup window. E-3 and H-1B1
# volumes are a small fraction of H-1B's, so the same bar would make every E-3
# employer look weak. Override per run with --strong.
DEFAULT_STRONG = {"e3": 5, "h1b": 25, "h1b1": 5}

LEGAL_SUFFIXES = {
    "inc", "incorporated", "llc", "llp", "lp", "ltd", "limited", "corp", "corporation",
    "co", "company", "plc", "pllc", "pc", "pa", "na", "sa", "ag", "gmbh", "bv", "pty",
    "the", "com", "dba",
}
TITLE_STOPWORDS = {"and", "of", "the", "for", "in", "a", "an", "i", "ii", "iii", "iv", "sr", "jr", "senior", "junior", "lead", "staff", "principal"}

DATE_FORMATS = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y", "%m/%d/%y", "%d-%b-%Y", "%Y/%m/%d")

WAGE_FACTORS = {"year": 1, "hour": 2080, "month": 12, "week": 52, "bi-weekly": 26, "biweekly": 26}


# ---------------------------------------------------------------------------
# Normalisation
# ---------------------------------------------------------------------------

def _ascii_fold(text: str) -> str:
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")


def normalize_employer(name) -> str:
    """Lower-case, strip accents, punctuation and legal suffixes.

    "Amazon.com Services LLC" -> "amazon services"; "ATLASSIAN, INC." -> "atlassian".
    Abbreviations written with dots ("L.L.C.") collapse before suffix removal.
    """
    if name is None:
        return ""
    text = _ascii_fold(str(name)).lower().replace("&", " and ")
    text = re.sub(r"\b([a-z])\.(?=[a-z]\.)", r"\1", text)  # l.l.c. -> llc.
    text = re.sub(r"\b([a-z])\.\b", r"\1", text)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    tokens = [t for t in text.split() if t not in LEGAL_SUFFIXES]
    return " ".join(tokens)


def normalize_title(title) -> str:
    text = _ascii_fold(str(title or "")).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def title_tokens(title: str) -> set[str]:
    return {t for t in normalize_title(title).split() if t not in TITLE_STOPWORDS and len(t) > 1}


def visa_key_for_lca(value) -> str | None:
    text = " ".join(str(value or "").lower().split())
    if not text:
        return None
    for key, spec in VISA_CLASSES.items():
        if text in spec["lca"] or (len(text) == 1 and text in spec["codes"]):
            return key
    return None


def status_bucket(value) -> str | None:
    text = re.sub(r"\s*-\s*", "-", " ".join(str(value or "").lower().split()))
    if text in ("certified", "certified-withdrawn"):
        return "certified"
    if text == "denied":
        return "denied"
    if text == "withdrawn":
        return "withdrawn"
    return None


def parse_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        # Excel serial day number (1900 date system).
        try:
            return date(1899, 12, 30) + timedelta(days=int(value))
        except (OverflowError, ValueError):
            return None
    text = str(value).strip()
    for candidate in (text, text[:19].strip(), text[:10].strip()):
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(candidate, fmt).date()
            except ValueError:
                continue
    match = re.match(r"(\d{4})-(\d{2})-(\d{2})", text)
    if match:
        try:
            return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
        except ValueError:
            return None
    return None


def fiscal_year(day: date) -> int:
    """US federal fiscal year: FY2026 runs 2025-10-01 to 2026-09-30."""
    return day.year + 1 if day.month >= 10 else day.year


def parse_number(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[^0-9.\-]", "", str(value))
    if not text or text in (".", "-"):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def annual_wage(amount, unit) -> float | None:
    number = parse_number(amount)
    if number is None or number <= 0:
        return None
    factor = WAGE_FACTORS.get(" ".join(str(unit or "year").lower().split()), None)
    if factor is None:
        return None
    wage = number * factor
    return wage if 1_000 <= wage <= 5_000_000 else None


# ---------------------------------------------------------------------------
# Reading source files
# ---------------------------------------------------------------------------

def _decode_bytes(raw: bytes) -> str:
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        return raw.decode("utf-16")
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8", "replace")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("cp1252", "replace")


def iter_rows(path: Path):
    """Yield (header list, row iterator) for a .csv/.tsv/.txt or .xlsx file."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            import openpyxl  # noqa: WPS433 - optional dependency
        except ImportError:
            sys.exit(
                f"{path.name}: reading .xlsx needs openpyxl (pip install openpyxl), "
                "or open the file and save it as CSV first."
            )
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            header = [str(h or "").strip() for h in next(rows)]
        except StopIteration:
            return [], iter(())
        return header, rows
    text = _decode_bytes(path.read_bytes())
    sample = text[:4096]
    delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    try:
        header = [h.strip() for h in next(reader)]
    except StopIteration:
        return [], iter(())
    return header, reader


def _key(header: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", header.upper()).strip("_")


def detect_kind(header: list[str]) -> str | None:
    keys = {_key(h) for h in header}
    if "VISA_CLASS" in keys and "CASE_STATUS" in keys:
        return "lca"
    lowered = [h.lower() for h in header]
    if any("employer" in h or "petitioner" in h for h in lowered) and any("approv" in h for h in lowered):
        return "uscis"
    return None


def _fy_from_name(path: Path) -> int | None:
    match = re.search(r"FY[\s_-]?(\d{4})", path.name, re.IGNORECASE)
    return int(match.group(1)) if match else None


class Builder:
    def __init__(self, classes: set[str]):
        self.classes = classes
        self.employers: dict[str, dict] = {}
        self.raw_names: dict[str, Counter] = defaultdict(Counter)
        self.titles: dict[tuple, Counter] = defaultdict(Counter)
        self.states: dict[tuple, Counter] = defaultdict(Counter)
        self.wages: dict[tuple, list] = defaultdict(list)
        self.seen_cases: set[str] = set()
        self.sources: list[dict] = []
        self.latest_decision: dict[str, str] = {}

    def _employer(self, raw_name) -> tuple[str, dict] | None:
        norm = normalize_employer(raw_name)
        if not norm:
            return None
        self.raw_names[norm][" ".join(str(raw_name).split())] += 1
        return norm, self.employers.setdefault(norm, {})

    def add_lca(self, path: Path, header: list[str], rows) -> dict:
        index = {_key(h): i for i, h in enumerate(header)}

        def col(*names):
            for name in names:
                if name in index:
                    return index[name]
            return None

        c_case = col("CASE_NUMBER", "LCA_CASE_NUMBER")
        c_status = col("CASE_STATUS", "STATUS")
        c_visa = col("VISA_CLASS")
        c_employer = col("EMPLOYER_NAME", "LCA_CASE_EMPLOYER_NAME")
        c_title = col("JOB_TITLE", "LCA_CASE_JOB_TITLE")
        c_state = col("WORKSITE_STATE", "WORKSITE_STATE_1", "LCA_CASE_WORKLOC1_STATE", "EMPLOYER_STATE")
        c_decision = col("DECISION_DATE", "CASE_DECISION_DATE", "DECISION_DATE_1")
        c_received = col("RECEIVED_DATE", "CASE_SUBMITTED", "SUBMIT_DATE")
        c_wage = col("WAGE_RATE_OF_PAY_FROM", "WAGE_RATE_OF_PAY_FROM_1", "LCA_CASE_WAGE_RATE_FROM", "WAGE_RATE_OF_PAY")
        c_unit = col("WAGE_UNIT_OF_PAY", "WAGE_UNIT_OF_PAY_1", "LCA_CASE_WAGE_RATE_UNIT")
        if c_status is None or c_visa is None or c_employer is None:
            sys.exit(f"{path.name}: LCA file is missing CASE_STATUS, VISA_CLASS or EMPLOYER_NAME columns")
        name_fy = _fy_from_name(path)

        def cell(row, i):
            return row[i] if i is not None and i < len(row) else None

        stats = Counter()
        fys = set()
        for row in rows:
            if not row:
                continue
            stats["rows"] += 1
            visa = visa_key_for_lca(cell(row, c_visa))
            if visa is None or visa not in self.classes:
                continue
            bucket = status_bucket(cell(row, c_status))
            if bucket is None:
                continue
            case = cell(row, c_case)
            if case:
                case = str(case).strip()
                if case in self.seen_cases:
                    stats["duplicates"] += 1
                    continue
                self.seen_cases.add(case)
            decided = parse_date(cell(row, c_decision)) or parse_date(cell(row, c_received))
            fy = fiscal_year(decided) if decided else name_fy
            if fy is None:
                stats["no_date"] += 1
                continue
            found = self._employer(cell(row, c_employer))
            if found is None:
                continue
            norm, record = found
            counts = record.setdefault("lca", {}).setdefault(visa, {}).setdefault(str(fy), {})
            counts[bucket] = counts.get(bucket, 0) + 1
            fys.add(fy)
            stats[visa] += 1
            if decided:
                key = visa
                if decided.isoformat() > self.latest_decision.get(key, ""):
                    self.latest_decision[key] = decided.isoformat()
            if bucket == "certified":
                title = normalize_title(cell(row, c_title))
                if title:
                    self.titles[(norm, visa)][title] += 1
                state = str(cell(row, c_state) or "").strip().upper()
                if state:
                    self.states[(norm, visa)][state] += 1
                wage = annual_wage(cell(row, c_wage), cell(row, c_unit))
                if wage:
                    self.wages[(norm, visa)].append(wage)
        source = {
            "file": path.name,
            "kind": "lca",
            "rows": stats["rows"],
            "kept": {k: stats[k] for k in self.classes if stats[k]},
            "duplicates_skipped": stats["duplicates"],
            "fiscal_years": sorted(fys),
        }
        self.sources.append(source)
        return source

    def add_uscis(self, path: Path, header: list[str], rows) -> dict:
        lowered = [h.lower() for h in header]
        c_employer = next(
            (i for i, h in enumerate(lowered) if ("employer" in h or "petitioner" in h) and "name" in h),
            next((i for i, h in enumerate(lowered) if h.strip() in ("employer", "petitioner")), None),
        )
        if c_employer is None:
            c_employer = next((i for i, h in enumerate(lowered) if "employer" in h or "petitioner" in h), None)
        c_fy = next((i for i, h in enumerate(lowered) if "fiscal" in h and "year" in h), None)
        c_approvals = [i for i, h in enumerate(lowered) if "approv" in h]
        c_denials = [i for i, h in enumerate(lowered) if "denial" in h or "denied" in h]
        if c_employer is None or not c_approvals:
            sys.exit(f"{path.name}: USCIS file needs an employer column and approval columns")
        name_fy = _fy_from_name(path)
        stats = Counter()
        fys = set()
        for row in rows:
            if not row or c_employer >= len(row):
                continue
            stats["rows"] += 1
            fy_value = parse_number(row[c_fy]) if c_fy is not None and c_fy < len(row) else None
            fy = int(fy_value) if fy_value else name_fy
            if fy is None:
                stats["no_date"] += 1
                continue
            found = self._employer(row[c_employer])
            if found is None:
                continue
            _, record = found
            approvals = sum(parse_number(row[i]) or 0 for i in c_approvals if i < len(row))
            denials = sum(parse_number(row[i]) or 0 for i in c_denials if i < len(row))
            counts = record.setdefault("uscis", {}).setdefault(str(fy), {"approvals": 0, "denials": 0})
            counts["approvals"] += int(approvals)
            counts["denials"] += int(denials)
            fys.add(fy)
        source = {"file": path.name, "kind": "uscis", "rows": stats["rows"], "fiscal_years": sorted(fys)}
        self.sources.append(source)
        return source

    def finish(self) -> dict:
        for norm, record in self.employers.items():
            names = self.raw_names[norm].most_common(3)
            record["name"] = names[0][0] if names else norm
            if len(names) > 1:
                record["aliases"] = [n for n, _ in names[1:]]
            for visa in list(record.get("lca", {})):
                key = (norm, visa)
                if self.titles.get(key):
                    record.setdefault("titles", {})[visa] = dict(self.titles[key].most_common(15))
                if self.states.get(key):
                    record.setdefault("states", {})[visa] = dict(self.states[key].most_common(15))
                if self.wages.get(key):
                    values = self.wages[key]
                    record.setdefault("wages", {})[visa] = {
                        "median_annual": int(statistics.median(values)),
                        "n": len(values),
                    }
        return {
            "metadata": {
                "version": DATA_VERSION,
                "built": date.today().isoformat(),
                "classes": sorted(self.classes),
                "latest_decision_date": self.latest_decision,
                "sources": self.sources,
            },
            "employers": self.employers,
        }


def cmd_import(args) -> int:
    classes = {c.strip().lower() for c in args.classes.split(",") if c.strip()}
    unknown = classes - set(VISA_CLASSES)
    if unknown:
        sys.exit(f"unknown visa class(es): {', '.join(sorted(unknown))}; expected {', '.join(VISA_CLASSES)}")
    builder = Builder(classes)
    for name in args.files:
        path = Path(name)
        if not path.is_file():
            sys.exit(f"not a file: {path}")
        header, rows = iter_rows(path)
        kind = args.kind if args.kind != "auto" else detect_kind(header)
        if kind == "lca":
            source = builder.add_lca(path, header, rows)
        elif kind == "uscis":
            if "h1b" not in classes:
                print(f"skipping {path.name}: USCIS data is H-1B only and h1b is not in --classes", file=sys.stderr)
                continue
            source = builder.add_uscis(path, header, rows)
        else:
            sys.exit(
                f"{path.name}: not recognised as a DOL LCA disclosure file (needs CASE_STATUS and VISA_CLASS "
                "columns) or a USCIS H-1B Employer Data Hub export (employer + approval columns). "
                "Pass --kind lca|uscis to force it."
            )
        print(f"read {path.name} ({source['kind']}): {source['rows']} rows, fiscal years {source['fiscal_years']}", file=sys.stderr)
    data = builder.finish()
    out = Path(args.out)
    tmp = out.with_suffix(out.suffix + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    tmp.replace(out)
    summary = {
        "written": str(out),
        "employers": len(data["employers"]),
        "classes": data["metadata"]["classes"],
        "sources": data["metadata"]["sources"],
    }
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    return 0


# ---------------------------------------------------------------------------
# Querying
# ---------------------------------------------------------------------------

def load_data(path: Path) -> dict:
    if not path.is_file():
        sys.exit(
            f"sponsor data not found at {path}. Download the DOL LCA disclosure files (and, for H-1B, "
            "the USCIS H-1B Employer Data Hub CSVs) and run: python3 tools/sponsor_lookup.py import <files> "
            "- see tools/README_SPONSOR_TOOL.md"
        )
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        sys.exit(f"cannot read {path}: {exc}")
    if not isinstance(data, dict) or not isinstance(data.get("employers"), dict):
        sys.exit(f"{path} is not a sponsor_lookup data file; rebuild it with the import command")
    return data


def available_years(data: dict, visa: str) -> list[int]:
    years = set()
    for record in data["employers"].values():
        for fy in record.get("lca", {}).get(visa, {}):
            years.add(int(fy))
        if visa == "h1b":
            for fy in record.get("uscis", {}):
                years.add(int(fy))
    return sorted(years)


def window_years(data: dict, visa: str, years: int) -> list[int]:
    have = available_years(data, visa)
    if not have:
        return []
    latest = have[-1]
    return [fy for fy in range(latest - years + 1, latest + 1)]


def partial_year_note(data: dict, visa: str, window: list[int]) -> str | None:
    latest = (data.get("metadata", {}).get("latest_decision_date") or {}).get(visa)
    if not latest or not window:
        return None
    try:
        day = date.fromisoformat(latest)
    except ValueError:
        return None
    if fiscal_year(day) == window[-1] and not (day.month == 9 and day.day >= 25):
        return f"FY{window[-1]} is partial: the newest filing in the data is decided {latest}."
    return None


def record_counts(record: dict, visa: str, window: list[int]) -> dict:
    lca = record.get("lca", {}).get(visa, {})
    per_year = {}
    certified = denied = withdrawn = 0
    for fy in window:
        counts = lca.get(str(fy), {})
        c = counts.get("certified", 0)
        certified += c
        denied += counts.get("denied", 0)
        withdrawn += counts.get("withdrawn", 0)
        if c:
            per_year[str(fy)] = c
    result = {"lca_certified": certified, "lca_denied": denied, "lca_withdrawn": withdrawn, "lca_by_year": per_year}
    if visa == "h1b":
        uscis = record.get("uscis", {})
        approvals = sum(uscis.get(str(fy), {}).get("approvals", 0) for fy in window)
        denials = sum(uscis.get(str(fy), {}).get("denials", 0) for fy in window)
        result["uscis_approvals"] = approvals
        result["uscis_denials"] = denials
    return result


def filings(counts: dict, visa: str) -> int:
    """The number the verdict is judged on: USCIS approvals for H-1B when the
    data has them (a petition decision), else certified LCAs."""
    if visa == "h1b" and counts.get("uscis_approvals"):
        return counts["uscis_approvals"]
    return counts["lca_certified"]


def find_matches(data: dict, query: str, exact_only: bool) -> tuple[str, list[str]]:
    q = normalize_employer(query)
    if not q:
        return "none", []
    employers = data["employers"]
    if q in employers:
        exact = [q]
        if exact_only:
            return "exact", exact
    else:
        exact = []
        if exact_only:
            return "none", []
    q_tokens = q.split()
    family = [k for k in employers if k != q and k.split()[: len(q_tokens)] == q_tokens]
    if exact or family:
        return ("exact" if exact else "family"), exact + family
    contains = [k for k in employers if set(q_tokens) <= set(k.split())]
    if contains:
        return "contains", contains
    return "none", []


def suggestions(data: dict, query: str) -> list[str]:
    import difflib

    q = normalize_employer(query)
    names = list(data["employers"])
    close = difflib.get_close_matches(q, names, n=5, cutoff=0.75)
    return [data["employers"][k].get("name", k) for k in close]


def title_match(record: dict, visa: str, title: str | None) -> dict | None:
    if not title:
        return None
    wanted = title_tokens(title)
    if not wanted:
        return None
    best = []
    for filed, count in (record.get("titles", {}).get(visa) or {}).items():
        have = title_tokens(filed)
        if not have:
            continue
        overlap = len(wanted & have) / len(wanted | have)
        if overlap >= 0.5 or wanted <= have:
            best.append((filed, count))
    best.sort(key=lambda x: -x[1])
    return {"matched": bool(best), "titles": [{"title": t, "count": c} for t, c in best[:5]]}


def verdict_for(total: int, strong: int) -> str:
    if total >= strong:
        return "Strong"
    if total >= 1:
        return "Some history"
    return "None found"


def sponsor_score(verdict: str, total: int, strong: int, visa: str, title_hit: bool) -> int:
    """0-100 Sponsor Strength score (04-job-evaluation.md). Absence of filings
    is weak evidence, not a refusal - least of all for E-3/H-1B1, which have
    no lottery and a cheap filing - so "None found" never scores zero."""
    if verdict == "Strong":
        score = 90
    elif verdict == "Some history":
        score = 55 + int(30 * min(total, strong) / max(strong, 1))
    else:
        score = 30 if visa in ("e3", "h1b1") else 15
    if title_hit and verdict != "None found":
        score += 10
    return max(0, min(100, score))


def describe_record(record: dict, visa: str, window: list[int], title: str | None) -> dict:
    counts = record_counts(record, visa, window)
    entity = {
        "name": record.get("name"),
        "filings": filings(counts, visa),
        **counts,
        "top_titles": list((record.get("titles", {}).get(visa) or {}).items())[:5],
        "top_states": list((record.get("states", {}).get(visa) or {}).items())[:5],
    }
    wage = (record.get("wages", {}) or {}).get(visa)
    if wage:
        entity["median_annual_wage"] = wage.get("median_annual")
    if record.get("aliases"):
        entity["aliases"] = record["aliases"]
    match = title_match(record, visa, title)
    if match is not None:
        entity["title_match"] = match
    return entity


def cmd_lookup(args) -> int:
    visa = args.visa.lower()
    if visa in NO_DATA_PATHS:
        result = {"query": args.company, "visa": visa, "verdict": "n/a", "score": None, "notes": [NO_DATA_PATHS[visa]]}
        return emit(result, args.json, text_lookup)
    if visa not in VISA_CLASSES:
        sys.exit(f"unknown --visa {visa!r}; expected one of {', '.join([*VISA_CLASSES, *NO_DATA_PATHS])}")
    data = load_data(Path(args.data))
    if visa not in data.get("metadata", {}).get("classes", []):
        sys.exit(f"the data file has no {visa} filings; re-run import with --classes including {visa}")
    window = window_years(data, visa, args.years)
    strong = args.strong if args.strong is not None else DEFAULT_STRONG[visa]
    mode, keys = find_matches(data, args.company, args.exact)
    query_key = normalize_employer(args.company)
    described = {k: describe_record(data["employers"][k], visa, window, args.title) for k in keys}
    entities = sorted(described.values(), key=lambda e: -e["filings"])
    notes = []
    if mode == "exact":
        basis = [described[query_key]]
        if len(entities) > 1:
            notes.append(
                f"{len(entities) - 1} related employer name(s) start with the same words; the verdict uses the "
                "exact name only - check the list in case a sister entity is the one hiring."
            )
    else:
        basis = entities
        if mode in ("family", "contains") and len(entities) > 1:
            notes.append(
                f"No exact name match: the verdict sums {len(entities)} employer names that {'start with' if mode == 'family' else 'contain'} "
                "the query - check the names, and re-run with the exact legal name (e.g. from the posting's footer) if one is wrong."
            )
    total = sum(e["filings"] for e in basis)
    title_hit = any((e.get("title_match") or {}).get("matched") for e in basis)
    verdict = verdict_for(total, strong)
    if verdict == "None found":
        notes.append(
            "No filings found under this name. That is weak evidence, not a refusal: the employer may file under "
            "a different legal name, may be new to sponsoring, or (for E-3/H-1B1) may simply never have had an "
            "Australian/Chilean/Singaporean hire yet."
        )
    partial = partial_year_note(data, visa, window)
    if partial:
        notes.append(partial)
    result = {
        "query": args.company,
        "visa": visa,
        "visa_label": VISA_CLASSES[visa]["label"],
        "window": [f"FY{fy}" for fy in window],
        "match": mode,
        "filings_in_window": total,
        "strong_threshold": strong,
        "verdict": verdict,
        "score": sponsor_score(verdict, total, strong, visa, title_hit),
        "meets_minimum": total >= args.min,
        "minimum": args.min,
        "title_match": title_hit if args.title else None,
        "entities": entities[: args.limit],
        "notes": notes,
    }
    if mode == "none":
        result["suggestions"] = suggestions(data, args.company)
    return emit(result, args.json, text_lookup)


def cmd_top(args) -> int:
    visa = args.visa.lower()
    if visa not in VISA_CLASSES:
        sys.exit(f"--visa must be one of {', '.join(VISA_CLASSES)} for top")
    data = load_data(Path(args.data))
    window = window_years(data, visa, args.years)
    state = args.state.upper() if args.state else None
    wanted = title_tokens(args.title) if args.title else set()
    rows = []
    for record in data["employers"].values():
        if visa not in record.get("lca", {}) and not (visa == "h1b" and record.get("uscis")):
            continue
        if state and state not in (record.get("states", {}).get(visa) or {}):
            continue
        titles = record.get("titles", {}).get(visa) or {}
        if wanted:
            hits = {t: c for t, c in titles.items() if wanted <= title_tokens(t) or len(wanted & title_tokens(t)) / max(len(wanted | title_tokens(t)), 1) >= 0.5}
            if not hits:
                continue
        counts = record_counts(record, visa, window)
        total = filings(counts, visa)
        if total <= 0:
            continue
        rows.append(
            {
                "name": record.get("name"),
                "filings": total,
                "top_titles": list(titles.items())[:3],
                "top_states": list((record.get("states", {}).get(visa) or {}).items())[:3],
            }
        )
    rows.sort(key=lambda r: (-r["filings"], r["name"] or ""))
    result = {
        "visa": visa,
        "visa_label": VISA_CLASSES[visa]["label"],
        "window": [f"FY{fy}" for fy in window],
        "state": state,
        "title": args.title,
        "employers": rows[: args.limit],
        "total_employers": len(rows),
    }
    note = partial_year_note(data, visa, window)
    if note:
        result["notes"] = [note]
    return emit(result, args.json, text_top)


def cmd_info(args) -> int:
    data = load_data(Path(args.data))
    meta = data.get("metadata", {})
    result = {
        "data_file": str(args.data),
        "built": meta.get("built"),
        "classes": meta.get("classes"),
        "employers": len(data["employers"]),
        "fiscal_years": {visa: [f"FY{y}" for y in available_years(data, visa)] for visa in meta.get("classes", [])},
        "latest_decision_date": meta.get("latest_decision_date"),
        "sources": meta.get("sources"),
    }
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0


# ---------------------------------------------------------------------------
# Posting text scan
# ---------------------------------------------------------------------------

NEGATION = re.compile(r"\b(no|not|never|unable|cannot|can't|won't|will not|do not|does not|don't|doesn't|without)\b[^.;:!?]{0,30}$", re.I)

CITIZEN_ONLY = [
    r"\b(u\.?\s?s\.?|united states)\s+citizens?\s+only\b",
    r"\bonly\s+(u\.?\s?s\.?|united states)\s+citizens\b",
    r"\bmust\s+be\s+(a\s+)?(u\.?\s?s\.?|united states)\s+citizen\b",
    r"\b(u\.?\s?s\.?|united states)\s+citizenship\s+(is\s+)?required\b",
    r"\brequires?\s+(u\.?\s?s\.?|united states)\s+citizenship\b",
    r"\b(u\.?\s?s\.?\s+)?citizens?\s+or\s+(lawful\s+)?permanent\s+residents?\s+only\b",
    r"\bgreen\s+card\s+holders?\s+only\b",
    r"\bonly\s+(u\.?\s?s\.?\s+)?(citizens|green\s+card\s+holders|permanent\s+residents)\b",
]
CLEARANCE = [
    r"\b(active|current)\s+(secret|top\s+secret|ts(/sci)?|public\s+trust)\b",
    r"\b(secret|top\s+secret|ts/sci|ts\s*/\s*sci|public\s+trust)\s+(security\s+)?clearance\b",
    r"\bsecurity\s+clearance\s+(is\s+)?required\b",
    r"\b(must|able\s+to)\s+(obtain|hold|maintain)\s+(a|an|the)?\s*(u\.?s\.?\s+)?(government\s+)?(security\s+)?clearance\b",
    r"\bclearance\s+(is\s+)?required\b",
]
NO_SPONSORSHIP = [
    r"\b(not|unable\s+to|cannot|can't|won't|will\s+not|does\s+not|do\s+not|don't|doesn't)\s+(currently\s+)?(be\s+)?(able\s+to\s+)?(provide|offer|support|consider)?\s*(any\s+)?(employment\s+)?(visa\s+|work\s+)?sponsor(ship|ing|s)?\b",
    r"\b(will\s+not|do\s+not|does\s+not|don't|doesn't|won't)\s+(now\s+or\s+in\s+the\s+future\s+)?(currently\s+)?(require|need)\s+(visa\s+|employer\s+|employment\s+)?sponsorship\b",
    r"\bno\s+(employment\s+)?(visa\s+|work\s+)?sponsorship\b",
    r"\bsponsorship\s+(is\s+|will\s+)?(not|un)\s*(be\s+)?(available|provided|offered|possible|considered)\b",
    r"\bnot\s+eligible\s+for\s+(visa\s+)?sponsorship\b",
    r"\bwithout\s+(the\s+need\s+for\s+)?(current\s+or\s+future\s+|now\s+or\s+in\s+the\s+future\s+)?(visa\s+|employer\s+)?sponsorship\b",
    r"\b(must|need\s+to)\s+(be\s+)?(legally\s+)?authori[sz]ed\s+to\s+work\s+in\s+the\s+(u\.?s\.?|united\s+states)\s+(now\s+and\s+in\s+the\s+future\s+)?without\b",
    r"\bsponsorship\s+(is\s+)?not\s+(an\s+option|offered)\b",
]
POSITIVE = [
    r"\b(visa\s+)?sponsorship\s+(is\s+)?(available|provided|offered|possible)\b",
    r"\b(will|can|do|does|happy\s+to|able\s+to|willing\s+to)\s+sponsor\b",
    r"\bwe\s+sponsor\b",
    r"\bsponsor(s|ing)?\s+(h-?1b|e-?3|visas?|work\s+visas?)\b",
    r"\bopen\s+to\s+(visa\s+)?sponsorship\b",
    r"\bh-?1b\s+transfers?\s+(welcome|accepted|considered)\b",
    r"\b(e-?3|h-?1b1?)\s+(visa\s+)?(holders?|candidates?)\s+(are\s+)?(welcome|encouraged|considered)\b",
]
# "Must be a US citizen OR authorized to work in the US" is a work-authorization
# requirement, not a citizens-only rule: the match is dropped when this follows it.
CITIZEN_OR_AUTHORIZED = re.compile(
    r"[^.;]{0,40}?\bor\s+(be\s+|have\s+|hold\s+)?(an?\s+)?(legally\s+|otherwise\s+)?(authori[sz]|eligible|permitted|work\s+authori[sz]ation|valid\s+work)",
    re.I,
)
SPONSORSHIP_QUESTION = re.compile(r"\b(will|do|would)\s+you\s+(now\s+or\s+in\s+the\s+future\s+)?(require|need)\s+(visa\s+|employer\s+)?sponsorship\b", re.I)
MENTIONS = {
    "e3": re.compile(r"\be-?3\s+(visa|status|sponsorship|holders?|candidates?)\b|\be-?3\s*\(australia", re.I),
    "h1b": re.compile(r"\bh-?1b\b(?!1)", re.I),
    "h1b1": re.compile(r"\bh-?1b1\b", re.I),
    "tn": re.compile(r"\btn\s+(visa|status)\b", re.I),
    "opt": re.compile(r"\b(stem\s+)?opt\b(?![-\s]?(in|out)\b)|\bcpt\b", re.I),
}


def _quote(text: str, start: int, end: int, pad: int = 70) -> str:
    snippet = text[max(0, start - pad): min(len(text), end + pad)]
    snippet = " ".join(snippet.split())
    prefix = "..." if start > pad else ""
    suffix = "..." if end + pad < len(text) else ""
    return f"{prefix}{snippet}{suffix}"


def _scan(patterns, text, *, skip_negated=False, skip_question=False, skip_or_authorized=False):
    hits = []
    for pattern in patterns:
        for match in re.finditer(pattern, text, re.I):
            if skip_or_authorized and CITIZEN_OR_AUTHORIZED.match(text, match.end()):
                continue
            if skip_negated and NEGATION.search(text[max(0, match.start() - 40): match.start()]):
                continue
            if skip_negated and re.match(r"\s+(is\s+|will\s+)?(not|un)", text[match.end(): match.end() + 12], re.I):
                continue
            if skip_question:
                around = text[max(0, match.start() - 60): match.end() + 10]
                if SPONSORSHIP_QUESTION.search(around):
                    continue
            hits.append(match)
    # One quote per overlapping region.
    hits.sort(key=lambda m: m.start())
    unique, last_end = [], -1
    for match in hits:
        if match.start() > last_end:
            unique.append(match)
            last_end = match.end() + 40
    return unique


def check_posting_text(text: str, *, visa: str, needs_sponsorship: bool, no_sponsorship: str, has_clearance: bool) -> dict:
    text = text or ""
    findings = []
    citizen = _scan(CITIZEN_ONLY, text, skip_or_authorized=True)
    clearance = _scan(CLEARANCE, text)
    negative = _scan(NO_SPONSORSHIP, text, skip_question=True)
    positive = _scan(POSITIVE, text, skip_negated=True, skip_question=True)
    # A positive phrase inside a negative one ("no visa sponsorship") is not positive.
    positive = [p for p in positive if not any(n.start() <= p.start() <= n.end() + 5 for n in negative)]
    for category, hits in (("citizen_only", citizen), ("clearance", clearance), ("no_sponsorship", negative), ("sponsorship_offered", positive)):
        for match in hits:
            findings.append({"category": category, "quote": _quote(text, match.start(), match.end())})
    asks = bool(SPONSORSHIP_QUESTION.search(text))
    mentions = {key: bool(rx.search(text)) for key, rx in MENTIONS.items()}

    notes = []
    if citizen:
        verdict = "FAIL"
        notes.append("Citizenship or permanent-residency requirement - hard stop (Eligibility Gate).")
    elif clearance and not has_clearance:
        verdict = "FAIL"
        notes.append("Security clearance required; US clearances are gated on citizenship.")
    elif negative and needs_sponsorship:
        h1b_only = all(
            MENTIONS["h1b"].search(f["quote"]) and not re.search(r"\b(any|all|other)\b", f["quote"], re.I)
            for f in findings if f["category"] == "no_sponsorship"
        )
        if no_sponsorship == "flag" or (visa in ("e3", "h1b1") and h1b_only):
            verdict = "FLAG"
            if visa in ("e3", "h1b1") and h1b_only:
                notes.append(
                    "The no-sponsorship line names H-1B only. E-3/H-1B1 is a different, cheaper filing with no "
                    "lottery - worth asking the recruiter before ruling it out."
                )
            else:
                notes.append("Posting says no sponsorship; your profile says flag rather than drop these.")
        else:
            verdict = "FAIL"
            notes.append("Posting says no sponsorship and you need it; your profile says drop these.")
    elif positive:
        verdict = "PASS"
        notes.append("Posting says sponsorship is available.")
    elif not needs_sponsorship:
        verdict = "PASS"
    else:
        verdict = "UNVERIFIED"
        notes.append("Posting is silent on sponsorship - silence is not permission; check the employer's careers page or ask.")
    if asks and needs_sponsorship:
        notes.append(
            'The application asks whether you will need sponsorship: answer it exactly as your profile records it.'
        )
    if mentions["e3"] and visa == "e3":
        notes.append("Posting mentions E-3 by name.")
    return {
        "verdict": verdict,
        "visa": visa,
        "needs_sponsorship": needs_sponsorship,
        "findings": findings,
        "asks_sponsorship_question": asks,
        "mentions": mentions,
        "notes": notes,
    }


def cmd_check_posting(args) -> int:
    if args.infile:
        text = Path(args.infile).read_text(encoding="utf-8", errors="replace")
    else:
        # Bytes, not sys.stdin.read(): a piped stdin on Windows decodes with the
        # ANSI code page and dies on the first non-cp1252 byte of a posting.
        text = sys.stdin.buffer.read().decode("utf-8", "replace")
    guard = None
    if text.lstrip().startswith("posting_guard:") and "\n--- text ---\n" in text:
        # posting_guard.py's combined output: scan only the cleaned text, so
        # hidden blocks the gate removed can neither pass nor fail the check.
        guard = text.lstrip().splitlines()[0].strip()
        text = text.split("\n--- text ---\n", 1)[1]
    result = check_posting_text(
        text,
        visa=args.visa.lower(),
        needs_sponsorship=args.needs_sponsorship == "yes",
        no_sponsorship=args.no_sponsorship,
        has_clearance=args.has_clearance,
    )
    if guard:
        result["guard"] = guard
    return emit(result, args.json, text_check)


# ---------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------

def emit(result: dict, as_json: bool, formatter) -> int:
    if as_json:
        print(json.dumps(result, indent=2, ensure_ascii=False))
    else:
        print(formatter(result))
    return 0


def text_lookup(r: dict) -> str:
    lines = [f"Sponsor lookup: {r['query']} - {r.get('visa_label', r['visa'])}"]
    if r.get("verdict") == "n/a":
        lines += [f"  {n}" for n in r.get("notes", [])]
        return "\n".join(lines)
    lines.append(f"  Window: {', '.join(r['window']) or 'no data'}   Match: {r['match']}")
    lines.append(f"  Verdict: {r['verdict']} ({r['filings_in_window']} filings; Strong >= {r['strong_threshold']})   Score: {r['score']}/100")
    lines.append(f"  Meets your minimum ({r['minimum']}): {'yes' if r['meets_minimum'] else 'no'}")
    for e in r.get("entities", []):
        lines.append(f"  - {e['name']}: {e['filings']} filings {e.get('lca_by_year') or ''}")
        if e.get("top_titles"):
            lines.append("      titles: " + "; ".join(f"{t} ({c})" for t, c in e["top_titles"]))
        if e.get("median_annual_wage"):
            lines.append(f"      median offered wage: ${e['median_annual_wage']:,}/yr")
    for s in r.get("suggestions", []) or []:
        lines.append(f"  did you mean: {s}")
    lines += [f"  note: {n}" for n in r.get("notes", [])]
    return "\n".join(lines)


def text_top(r: dict) -> str:
    lines = [f"Top {r['visa_label']} filers, {', '.join(r['window'])}" + (f", state {r['state']}" if r["state"] else "") + (f", title ~ {r['title']!r}" if r["title"] else "")]
    for i, e in enumerate(r["employers"], 1):
        titles = "; ".join(t for t, _ in e["top_titles"])
        lines.append(f"{i:>3}. {e['name']} - {e['filings']} filings - {titles}")
    lines.append(f"({r['total_employers']} employers matched)")
    lines += [f"note: {n}" for n in r.get("notes", [])]
    return "\n".join(lines)


def text_check(r: dict) -> str:
    lines = [f"Sponsorship check: {r['verdict']}"]
    for f in r["findings"]:
        lines.append(f"  [{f['category']}] \"{f['quote']}\"")
    lines += [f"  note: {n}" for n in r["notes"]]
    return "\n".join(lines)


def _force_utf8_output() -> None:
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure:
            reconfigure(encoding="utf-8")


def main(argv=None) -> int:
    _force_utf8_output()
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = ap.add_subparsers(dest="command", required=True)
    data_arg = argparse.ArgumentParser(add_help=False)
    data_arg.add_argument("--data", default=str(DATA_FILE), help="sponsor data file (default: repo-root sponsor_data.json)")

    imp = sub.add_parser("import", help="build the data file from downloaded DOL/USCIS files")
    imp.add_argument("files", nargs="+")
    imp.add_argument("--classes", default="e3,h1b,h1b1", help="visa classes to keep (default e3,h1b,h1b1)")
    imp.add_argument("--kind", choices=("auto", "lca", "uscis"), default="auto")
    imp.add_argument("--out", default=str(DATA_FILE))
    imp.set_defaults(func=cmd_import)

    info = sub.add_parser("info", parents=[data_arg], help="what the data file covers")
    info.set_defaults(func=cmd_info)

    look = sub.add_parser("lookup", parents=[data_arg], help="one employer's record for a visa class")
    look.add_argument("company")
    look.add_argument("--visa", required=True, help="e3 | h1b | h1b1 (or tn / o1 / l1 / none: no public data)")
    look.add_argument("--title", help="the posting's job title, to check for same-role filings")
    look.add_argument("--years", type=int, default=3, help="fiscal-year window ending at the newest year in the data")
    look.add_argument("--min", type=int, default=1, help="your minimum record (filings in the window) for strict mode")
    look.add_argument("--strong", type=int, help="filings needed for a Strong verdict (default: e3 5, h1b 25, h1b1 5)")
    look.add_argument("--exact", action="store_true", help="exact normalised name only")
    look.add_argument("--limit", type=int, default=8, help="max employer names to list")
    look.add_argument("--json", action="store_true")
    look.set_defaults(func=cmd_lookup)

    top = sub.add_parser("top", parents=[data_arg], help="employers with the most filings for a visa class")
    top.add_argument("--visa", required=True)
    top.add_argument("--state", help="two-letter worksite state, e.g. NY")
    top.add_argument("--title", help="keep employers who filed for a similar title")
    top.add_argument("--years", type=int, default=3)
    top.add_argument("--limit", type=int, default=25)
    top.add_argument("--json", action="store_true")
    top.set_defaults(func=cmd_top)

    chk = sub.add_parser("check-posting", help="classify a posting's sponsorship language (text on stdin or --in)")
    chk.add_argument("--in", dest="infile")
    chk.add_argument("--visa", default="h1b")
    chk.add_argument("--needs-sponsorship", choices=("yes", "no"), default="yes")
    chk.add_argument("--no-sponsorship", choices=("drop", "flag"), default="drop", help="what your profile says to do with 'no sponsorship' postings")
    chk.add_argument("--has-clearance", action="store_true")
    chk.add_argument("--json", action="store_true")
    chk.set_defaults(func=cmd_check_posting)

    args = ap.parse_args(argv)
    if getattr(args, "years", 1) < 1:
        sys.exit("--years must be at least 1")
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
